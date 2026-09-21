"""코스피200 데이터 읽기 · 계산 · 내보내기 (화면과 무관한 부분)

이 모듈은 인터넷에 접속하지 않는다. 한국거래소(KRX) 정보데이터시스템(https://data.krx.co.kr)에서
사람이 직접 내려받은 KOSPI200 구성종목/시세 파일을 읽는다.

읽을 수 있는 형식 (확장자가 아니라 파일 내용으로 판별한다)
    - CSV / 탭 구분 텍스트   (UTF-8, UTF-8 BOM, CP949 = KRX 기본)
    - Excel .xlsx            (openpyxl 필요)
    - HTML 표로 저장된 '엑셀' 파일  (BeautifulSoup4 로 표를 읽는다)
    옛 엑셀 .xls(BIFF)는 지원하지 않는다. 엑셀에서 .xlsx 로 저장하거나 CSV 로 다시 받는다.

열 이름은 별칭으로 찾는다 (KRX 마다 조금씩 다른 이름을 모두 인정): 종목코드 / 종목명 / 종가(현재가) / 대비 /
등락률 / 시가 / 고가 / 저가 / 거래량 / 거래대금 / 시가총액 / 상장주식수 / 시장구분 / 소속부.
숫자에 붙은 쉼표, %, 원, ▲▼, 괄호 음수는 자동으로 해석한다.
"""

import csv
import io
import math
import os
import re
from dataclasses import dataclass, field

from bs4 import BeautifulSoup

KOSPI200_SIZE = 200

# 화면/내보내기에서 쓰는 열 순서
COLUMNS = ("rank", "code", "name", "market", "sector", "close", "change", "rate",
           "open", "high", "low", "volume", "value", "marcap", "shares")
LABELS = {
    "rank": "순위", "code": "종목코드", "name": "종목명", "market": "시장구분", "sector": "소속부", "close": "종가",
    "change": "대비", "rate": "등락률(%)", "open": "시가", "high": "고가", "low": "저가",
    "volume": "거래량", "value": "거래대금", "marcap": "시가총액", "shares": "상장주식수",
}
TEXT_KEYS = ("code", "name", "market", "sector")
NUMERIC_KEYS = tuple(k for k in COLUMNS if k not in TEXT_KEYS)
ALIASES = {
    "rank": ("순위", "등수"),
    "code": ("종목코드", "단축코드", "종목번호", "코드"),
    "name": ("종목명", "한글종목약명", "한글종목명", "종목약명", "종목"),
    "market": ("시장구분", "시장"),
    "sector": ("소속부", "업종명", "업종"),
    "close": ("종가", "현재가", "종가(원)", "종가(현재가)"),
    "change": ("대비", "전일대비", "대비(원)"),
    "rate": ("등락률", "등락률(%)", "등락율", "등락율(%)"),
    "open": ("시가", "시가(원)"),
    "high": ("고가", "고가(원)"),
    "low": ("저가", "저가(원)"),
    "volume": ("거래량", "거래량(주)"),
    "value": ("거래대금", "거래대금(원)"),
    "marcap": ("시가총액", "시가총액(원)", "시가총액(백만원)", "시가총액(억원)"),
    "shares": ("상장주식수", "상장주식수(주)", "상장주식"),
}
TOTAL_ROW_NAMES = ("합계", "총계", "전체", "계", "total", "sum")
EMPTY_MARKS = ("", "-", "--", "—", "n/a", "na", "nan", "null", "none")
HEADER_SEARCH_ROWS = 15
CODE_RE = re.compile(r"[0-9A-Z]{6}")          # 종목코드: 숫자 6자리 (일부 종목은 끝이 영문자, 예: 00104K)


class KospiFileError(Exception):
    """파일을 읽을 수 없거나 코스피200 표로 볼 수 없을 때. 메시지는 그대로 사용자에게 보여 준다."""


@dataclass
class LoadResult:
    records: list                       # dict 목록 (COLUMNS 의 키)
    columns: list                       # 실제로 찾은 열 키 (COLUMNS 순서)
    skipped: int = 0                    # 종목이 아니라서 건너뛴 줄 수
    fmt: str = ""                       # 읽은 형식 설명
    notes: list = field(default_factory=list)
    source: str = ""                    # 파일 이름


# ---------------------------------------------------------------- 값 해석
def normalize_header(text):
    return re.sub(r"\s+", "", str(text if text is not None else "")).lower()


def parse_number(value):
    """'1,234' '+1.5%' '▼500' '(300)' 같은 값을 숫자로. 비어 있거나 해석할 수 없으면 None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return int(value) if float(value).is_integer() else float(value)

    s = str(value).strip()
    if s.lower() in EMPTY_MARKS:
        return None
    sign = 1
    if s[0] in "▲△":
        s = s[1:]
    elif s[0] in "▼▽":
        sign, s = -1, s[1:]
    if s.startswith("(") and s.endswith(")"):
        sign, s = -sign, s[1:-1]
    s = s.replace("−", "-").replace("－", "-").replace("＋", "+")
    s = re.sub(r"[,\s%원주]", "", s)
    if s.startswith("+"):
        s = s[1:]
    try:
        number = float(s) * sign
    except ValueError:
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return int(number) if number.is_integer() else number


def normalize_code(value):
    """엑셀이 앞자리 0 을 지운 코드(5930)도 6자리(005930)로 복원한다. 'A005930' 형태는 A 를 뗀다."""
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value)).zfill(6)
    s = str(value).strip().upper()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    if re.fullmatch(r"A\d{6}", s):
        s = s[1:]
    if s.isdigit() and len(s) < 6:
        s = s.zfill(6)
    return s


def _text(value):
    return "" if value is None else str(value).strip()


# ---------------------------------------------------------------- 파일 읽기 (원시 표 -> 줄 목록)
def _decode(raw):
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc), ("utf-8" if enc.startswith("utf") else "cp949")
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(일부 글자 깨짐)"


def _read_delimited(text):
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = "\t" if first.count("\t") > first.count(",") else ","
    try:
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as e:
        raise KospiFileError(f"CSV/텍스트로 읽을 수 없는 파일입니다 ({e}). KRX 에서 내려받은 CSV 또는 Excel 파일인지 확인하세요.") from e


def _read_html(raw):
    text, _ = _decode(raw)
    soup = BeautifulSoup(text, "html.parser")
    tables = soup.find_all("table")
    if not tables:
        raise KospiFileError("HTML 파일에서 표(<table>)를 찾지 못했습니다.")
    table = max(tables, key=lambda t: len(t.find_all("tr")))
    return [[cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])] for tr in table.find_all("tr")]


def _read_xlsx_sheets(path):
    """워크시트마다 (이름, 줄 목록) 을 돌려준다."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise KospiFileError("엑셀(.xlsx) 파일을 읽으려면 openpyxl 이 필요합니다: pip install openpyxl") from e
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:                                    # 손상된 zip, 다른 종류의 zip 등
        raise KospiFileError(f"엑셀 파일을 열 수 없습니다: {e.__class__.__name__}: {e}") from e
    sheets = []
    try:
        for ws in wb.worksheets:
            sheets.append((ws.title, [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]))
    finally:
        wb.close()
    return sheets


def read_raw_tables(path):
    """파일 내용을 보고 형식을 판별해 [(설명, 줄 목록), ...] 을 돌려준다."""
    try:
        with open(path, "rb") as f:
            raw = f.read()
    except OSError as e:
        raise KospiFileError(f"파일을 열 수 없습니다: {e}") from e
    if not raw.strip():
        raise KospiFileError("빈 파일입니다.")

    head = raw[:4096].lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if raw.startswith(b"PK"):
        return [(f"Excel 시트 '{name}'", rows) for name, rows in _read_xlsx_sheets(path)]
    if raw.startswith(b"\xd0\xcf\x11\xe0"):
        raise KospiFileError("옛 엑셀 형식(.xls)은 지원하지 않습니다. 엑셀에서 .xlsx 로 저장하거나 KRX 에서 CSV 로 다시 내려받으세요.")
    if head.startswith(b"<") or b"<table" in head:
        return [("HTML 표", _read_html(raw))]
    text, enc = _decode(raw)
    return [(f"CSV/텍스트 ({enc})", _read_delimited(text))]


# ---------------------------------------------------------------- 표 -> 종목 기록
def build_column_map(header_row):
    """머리글 한 줄에서 {열 키: 열 번호} 를 만든다. 알아볼 수 없는 열은 무시한다."""
    lookup = {normalize_header(a): key for key, names in ALIASES.items() for a in names}
    mapping = {}
    for idx, cell in enumerate(header_row):
        key = lookup.get(normalize_header(cell))
        if key and key not in mapping:
            mapping[key] = idx
    return mapping


def find_header(rows):
    for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        mapping = build_column_map(row)
        if ("name" in mapping or "code" in mapping) and len(mapping) >= 3:
            return i, mapping
    return None


def direction(record):
    """'up' / 'down' / 'flat' / None(알 수 없음). 등락률을 우선하고 없으면 대비로 판단한다."""
    for key in ("rate", "change"):
        v = record.get(key)
        if v is not None:
            return "up" if v > 0 else "down" if v < 0 else "flat"
    return None


def _finish_record(rec, notes_flags):
    close, change, rate = rec.get("close"), rec.get("change"), rec.get("rate")
    if rate is None and close is not None and change is not None and close - change:
        rec["rate"] = round(change / (close - change) * 100, 2)          # 등락률 열이 없으면 종가와 대비로 계산
        notes_flags["rate_computed"] = True
    rate, change = rec.get("rate"), rec.get("change")
    if rate and change and (rate > 0) != (change > 0):
        rec["change"] = math.copysign(abs(change), rate)                # 대비가 부호 없이 저장된 파일 보정
        rec["change"] = int(rec["change"]) if float(rec["change"]).is_integer() else rec["change"]
        notes_flags["sign_fixed"] = True
    return rec


def rows_to_records(rows):
    found = find_header(rows)
    if found is None:
        seen = next((r for r in rows if any(_text(c) for c in r)), [])
        raise KospiFileError(
            "코스피200 표의 열(종목코드, 종목명, 종가 …)을 찾지 못했습니다.\n"
            f"파일 첫 줄의 열 이름: {[_text(c) for c in seen][:12]}\n"
            "KRX 에서 KOSPI200 구성종목/시세를 CSV 또는 Excel 로 내려받은 파일인지 확인하세요.")
    header_idx, mapping = found

    records, skipped, flags = [], 0, {}
    for row in rows[header_idx + 1:]:
        if not any(_text(c) for c in row):
            continue
        rec = {}
        for key, idx in mapping.items():
            cell = row[idx] if idx < len(row) else ""
            if key == "code":
                rec[key] = normalize_code(cell)
            elif key in TEXT_KEYS:
                rec[key] = _text(cell)
            else:
                rec[key] = parse_number(cell)
        code, name = rec.get("code", ""), rec.get("name", "")
        if "code" in mapping and not CODE_RE.fullmatch(code):
            skipped += 1                                                 # 합계/각주 줄: 종목코드(6자리)가 아니다
            continue
        if not name and not code:
            skipped += 1
            continue
        if name.lower() in TOTAL_ROW_NAMES:
            skipped += 1
            continue
        if not name:
            rec["name"] = code
        records.append(_finish_record(rec, flags))
    columns = [k for k in COLUMNS if k in mapping]
    if "name" not in columns:
        columns.insert(1 if "code" in columns else 0, "name")
        for rec in records:
            rec.setdefault("name", rec.get("code", ""))
    if not records:
        hint = " (종목코드가 6자리 형식이 아니라서 모두 제외되었습니다)" if skipped else ""
        raise KospiFileError(f"머리글은 찾았지만 종목 데이터가 한 줄도 없습니다{hint}.")
    return records, columns, skipped, flags


def load_file(path):
    """파일을 읽어 LoadResult 를 돌려준다. 읽을 수 없으면 KospiFileError."""
    tables = read_raw_tables(path)
    last_error = None
    for desc, rows in tables:
        try:
            records, columns, skipped, flags = rows_to_records(rows)
        except KospiFileError as e:
            last_error = e
            continue
        notes = []
        if len(records) != KOSPI200_SIZE:
            notes.append(f"KOSPI200 구성종목은 {KOSPI200_SIZE}개인데 이 파일에는 {len(records)}개가 있습니다 "
                         "(다른 지수나 전종목 파일이거나 일부만 내려받았을 수 있습니다).")
        if skipped:
            notes.append(f"종목이 아닌 줄 {skipped}개(합계·각주 등)는 건너뛰었습니다.")
        if flags.get("rate_computed"):
            notes.append("등락률 열이 없어 종가와 대비로 계산했습니다.")
        if flags.get("sign_fixed"):
            notes.append("대비에 부호가 없어 등락률의 부호에 맞춰 보정했습니다.")
        for missing in ("close", "rate"):
            if missing not in columns:
                notes.append(f"'{LABELS[missing]}' 열이 없습니다.")
        return LoadResult(records, columns, skipped, desc, notes, os.path.basename(path))
    raise last_error or KospiFileError("읽을 수 있는 표가 없습니다.")


# ---------------------------------------------------------------- 요약
def summarize(records):
    up = down = flat = unknown = 0
    rates, marcaps = [], []
    top_up = top_down = None
    for r in records:
        d = direction(r)
        if d == "up":
            up += 1
        elif d == "down":
            down += 1
        elif d == "flat":
            flat += 1
        else:
            unknown += 1
        if r.get("rate") is not None:
            rates.append(r["rate"])
            if r["rate"] > 0 and (top_up is None or r["rate"] > top_up["rate"]):
                top_up = r
            if r["rate"] < 0 and (top_down is None or r["rate"] < top_down["rate"]):
                top_down = r
        if r.get("marcap") is not None:
            marcaps.append(r["marcap"])
    return {
        "total": len(records), "up": up, "down": down, "flat": flat, "unknown": unknown,
        "avg_rate": sum(rates) / len(rates) if rates else None,
        "marcap_sum": sum(marcaps) if marcaps else None,
        "top_up": top_up, "top_down": top_down,
    }


# ---------------------------------------------------------------- 표시용 서식
def fmt_number(value):
    if value is None:
        return ""
    if isinstance(value, float) and not value.is_integer():
        return f"{value:,.2f}"
    return f"{int(value):,}"


def fmt_signed(value):
    if value is None:
        return ""
    return ("+" if value > 0 else "") + fmt_number(value)


def fmt_rate(value):
    if value is None:
        return ""
    return f"{value:+.2f}%" if value else "0.00%"


def format_cell(key, value):
    if key in TEXT_KEYS:
        return value or ""
    if key == "change":
        return fmt_signed(value)
    if key == "rate":
        return fmt_rate(value)
    return fmt_number(value)


# ---------------------------------------------------------------- 내보내기
NUMBER_FORMATS = {"rate": "0.00", "change": "#,##0;-#,##0;0"}
UP_COLOR, DOWN_COLOR = "D92D20", "2E6BE6"      # 엑셀 글자색 (상승 빨강, 하락 파랑: 한국식)
EXCEL_WIDTHS = {"rank": 7, "code": 11, "name": 22, "market": 10, "sector": 12, "close": 12, "change": 11, "rate": 11,
                "open": 12, "high": 12, "low": 12, "volume": 15, "value": 18, "marcap": 20, "shares": 16}


def export_xlsx(records, columns, path):
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RuntimeError("엑셀(.xlsx) 저장에는 openpyxl 이 필요합니다: pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "KOSPI200"
    for c, key in enumerate(columns, 1):
        cell = ws.cell(row=1, column=c, value=LABELS[key])
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = EXCEL_WIDTHS.get(key, 14)

    for r, rec in enumerate(records, 2):
        d = direction(rec)
        for c, key in enumerate(columns, 1):
            value = rec.get(key)
            if key in TEXT_KEYS:
                cell = ws.cell(row=r, column=c, value=ILLEGAL_CHARACTERS_RE.sub("", value or ""))
                cell.data_type = "s"           # 종목코드의 앞자리 0 유지 + '=' 로 시작해도 수식으로 실행되지 않음
                cell.alignment = Alignment(horizontal="center" if key == "code" else "left")
            else:
                cell = ws.cell(row=r, column=c, value=value)
                cell.number_format = NUMBER_FORMATS.get(key, "#,##0")
                cell.alignment = Alignment(horizontal="right")
                if key in ("change", "rate") and d in ("up", "down"):
                    cell.font = Font(color=UP_COLOR if d == "up" else DOWN_COLOR)
    # 스크롤해도 열 이름과 종목명(그 앞의 순위·코드 포함)이 보이도록 고정
    ws.freeze_panes = f"{get_column_letter(columns.index('name') + 2)}2" if "name" in columns else "A2"
    if records:
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def export_csv(records, columns, path):
    """엑셀에서 한글이 깨지지 않도록 utf-8-sig. 숫자는 쉼표 없이 원래 값 그대로 쓴다."""
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([LABELS[k] for k in columns])
        for rec in records:
            writer.writerow(["" if rec.get(k) is None else rec[k] for k in columns])


def export(records, columns, path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        export_xlsx(records, columns, path)
    elif ext == ".csv":
        export_csv(records, columns, path)
    elif ext == ".xls":
        raise RuntimeError("옛 엑셀 형식(.xls)은 지원하지 않습니다. 확장자를 .xlsx 로 바꿔 저장하세요.")
    else:
        raise RuntimeError("저장 형식은 .xlsx 또는 .csv 만 지원합니다.")


def save_error_text(error):
    if isinstance(error, PermissionError):
        return f"파일에 쓸 수 없습니다. 엑셀 등에서 같은 파일이 열려 있으면 닫고 다시 저장하세요. ({error})"
    return str(error)
